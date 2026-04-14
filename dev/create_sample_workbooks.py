"""
Create minimal workbooks under data/ for local --verify / --dry-run smoke tests.
Replace with real files before production use.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    data = root / "data"
    data.mkdir(parents=True, exist_ok=True)

    patcher = data / "Placement_Fee_Patching.xlsx"
    tracker = data / "Master_Tracker_Points_Average.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Underwriting"
    ws.append(["API Loan ID", "Origination Fee", "Address"])
    ws.append([999001, 1500.0, "123 Sample St"])
    tab = Table(
        displayName="tbl_of",
        ref="A1:C2",
    )
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False
    )
    ws.add_table(tab)
    wb.save(patcher)
    print("Wrote", patcher)

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "MA API ID"
    ws2["A1"] = "Label"
    ws2["B1"] = "API Loan ID"
    ws2["C1"] = "Creation Date"
    ws2["D1"] = "StaticCol"
    ws2["E1"] = "Origination Fee"
    ws2["B2"] = 999001
    ws2["C2"] = datetime(2026, 1, 10)
    ws2["E2"] = None
    wb2.save(tracker)
    print("Wrote", tracker)


if __name__ == "__main__":
    main()
